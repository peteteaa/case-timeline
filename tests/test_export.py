"""Exports: CSV/XLSX of the list (columns in the given order), SVG/PNG of the visual view."""

from __future__ import annotations

import csv
import io
import xml.etree.ElementTree as ET
from datetime import date
from pathlib import Path

import pytest

from pipeline import export, llm
from pipeline.schema import AGENT, Fact, SpanIn
from pipeline.store import Store

CASE = "c1"
SVG_NS = "{http://www.w3.org/2000/svg}"


def _embed(texts: list[str]) -> list[list[float]]:
    return [llm.hashed_embedding(t) for t in texts]


@pytest.fixture
def store(tmp_path: Path) -> Store:
    s = Store(tmp_path / "case.db", embedder=_embed)
    yield s
    s.close()


@pytest.fixture
def facts(store: Store) -> list[Fact]:
    text = "The complaint alleges that Acme hid the losses from investors."
    doc, _ = store.register_document(
        CASE,
        filename="complaint.pdf",
        rel_path="complaint.pdf",
        sha256="s1",
        location="evidence",
        kind="pleading",
    )
    store.save_parsed_text(doc.doc_id, text)
    quote = "Acme hid the losses"
    start = text.index(quote)
    span = SpanIn(doc_id=doc.doc_id, char_start=start, char_end=start + len(quote), quote=quote)

    issue = store.add_tag(CASE, "Fraud", by="alice")
    other = store.add_tag(CASE, "Damages", by="alice")
    smith = store.add_witness(CASE, "John Smith", by="alice")
    jones = store.add_witness(CASE, "Mary Jones", by="alice")

    key = store.add_fact(
        CASE,
        "Board meets <Acme> & votes",
        by="alice",
        date_start=date(2022, 3, 1),
        statuses=["key_fact", "disputed"],
        tag_ids=[issue.tag_id, other.tag_id],
        witness_ids=[smith.witness_id, jones.witness_id],
    )
    store.add_note(key.fact_id, author="alice", body="check minutes")
    store.add_note(key.fact_id, author="bob", body="second note")
    rng = store.add_fact(
        CASE,
        "Losses concealed",
        by=AGENT,
        date_start=date(2022, 3, 1),
        date_end=date(2022, 3, 31),
        date_certainty="approximate",
        date_text="March 2022",
        spans=[span],
    )
    rel = store.add_fact(
        CASE, '=HYPERLINK("http://evil")', by="alice", date_text="two weeks after the meeting"
    )
    undated = store.add_fact(CASE, 'Undated "quoted" fact', by="alice")
    later = store.add_fact(CASE, "Settlement talks", by="alice", date_start=date(2022, 9, 15))
    return [store.get_fact(f.fact_id) for f in (key, rng, later, rel, undated)]  # type: ignore[misc]


# --------------------------------------------------------------------------- cells


def test_format_date() -> None:
    def fact(**kw: object) -> Fact:
        base = dict(
            fact_id="F-1",
            case_id=CASE,
            title="t",
            origin="user",
            review="accepted",
            created_by="a",
            created_at="2024-01-01T00:00:00+00:00",
            updated_by="a",
            updated_at="2024-01-01T00:00:00+00:00",
        )
        return Fact.model_validate({**base, **kw})

    assert export.format_date(fact()) == "Undated"
    assert export.format_date(fact(date_text="the next day")) == "the next day"
    assert export.format_date(fact(date_start="2022-03-01")) == "2022-03-01"
    assert export.format_date(fact(date_start="2022-03-01", date_end="2022-03-01")) == "2022-03-01"
    assert (
        export.format_date(
            fact(date_start="2022-03-01", date_end="2022-03-31", date_certainty="approximate")
        )
        == "2022-03-01 – 2022-03-31 (approx.)"
    )


# --------------------------------------------------------------------------- CSV


def _csv_rows(data: bytes) -> list[list[str]]:
    assert data.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


def test_csv_default_columns(store: Store, facts: list[Fact]) -> None:
    rows = _csv_rows(export.to_csv(store, CASE, facts))
    assert rows[0] == list(export.COLUMNS.values())
    assert len(rows) == 1 + len(facts)
    key = dict(zip(rows[0], rows[1], strict=True))
    assert key == {
        "Date": "2022-03-01",
        "Title": "Board meets <Acme> & votes",
        "Statuses": "Key fact; Disputed",
        "Review": "Accepted",
        "Evidence": "0",
        "Tags": "Damages; Fraud",
        "Witnesses": "John Smith; Mary Jones",
        "Notes": "2",
    }
    rng = dict(zip(rows[0], rows[2], strict=True))
    assert rng["Date"] == "2022-03-01 – 2022-03-31 (approx.)"
    assert rng["Title"] == "[ALLEGED] Losses concealed"  # only pleading evidence
    assert rng["Review"] == "Draft"
    assert rng["Evidence"] == "1"
    # caller's order is kept: dated facts first, then the two undated ones
    assert [r[0] for r in rows[1:]] == [
        "2022-03-01",
        "2022-03-01 – 2022-03-31 (approx.)",
        "2022-09-15",
        "two weeks after the meeting",
        "Undated",
    ]


def test_csv_respects_column_order_and_subset(store: Store, facts: list[Fact]) -> None:
    rows = _csv_rows(export.to_csv(store, CASE, facts[:2], columns=["notes", "title", "date"]))
    assert rows == [
        ["Notes", "Title", "Date"],
        ["2", "Board meets <Acme> & votes", "2022-03-01"],
        ["0", "[ALLEGED] Losses concealed", "2022-03-01 – 2022-03-31 (approx.)"],
    ]


def test_csv_neutralizes_formulas_and_keeps_quotes(store: Store, facts: list[Fact]) -> None:
    rows = _csv_rows(export.to_csv(store, CASE, facts, columns=["title"]))
    titles = [r[0] for r in rows[1:]]
    assert '\'=HYPERLINK("http://evil")' in titles
    assert 'Undated "quoted" fact' in titles


def test_unknown_or_empty_columns_raise(store: Store, facts: list[Fact]) -> None:
    with pytest.raises(ValueError, match="bogus"):
        export.to_csv(store, CASE, facts, columns=["title", "bogus"])
    with pytest.raises(ValueError):
        export.to_xlsx(store, CASE, facts, columns=["confidence"])
    with pytest.raises(ValueError):
        export.to_csv(store, CASE, facts, columns=[])


def test_csv_empty(store: Store) -> None:
    assert _csv_rows(export.to_csv(store, CASE, [])) == [list(export.COLUMNS.values())]


# --------------------------------------------------------------------------- XLSX


def test_xlsx(store: Store, facts: list[Fact]) -> None:
    from openpyxl import load_workbook

    cols = ["title", "date", "statuses", "evidence"]
    wb = load_workbook(io.BytesIO(export.to_xlsx(store, CASE, facts, columns=cols)))
    ws = wb.active
    assert ws.freeze_panes == "A2"
    assert [c.value for c in ws[1]] == ["Title", "Date", "Statuses", "Evidence"]
    assert all(c.font.bold for c in ws[1])
    assert ws.max_row == 1 + len(facts)
    assert ws["A2"].value == "Board meets <Acme> & votes"
    assert ws["D3"].value == 1  # counts are numbers
    # key fact row highlighted, others not
    assert ws["A2"].fill.fgColor.rgb.endswith(export.KEY_FACT_FILL)
    assert ws["A3"].fill.fill_type is None
    # formula-looking titles stay text
    formula = ws["A5"]
    assert formula.value == '=HYPERLINK("http://evil")'
    assert formula.data_type == "s"
    assert ws.column_dimensions["A"].width >= 20


def test_xlsx_empty(store: Store) -> None:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(export.to_xlsx(store, CASE, []))).active
    assert [c.value for c in ws[1]] == list(export.COLUMNS.values())
    assert ws.max_row == 1


# --------------------------------------------------------------------------- SVG / PNG


def _groups(root: ET.Element, cls: str) -> list[ET.Element]:
    return [g for g in root.iter(f"{SVG_NS}g") if cls in (g.get("class") or "").split()]


def _texts(el: ET.Element) -> list[str]:
    return [t.text or "" for t in el.iter(f"{SVG_NS}text")]


@pytest.mark.parametrize("zoom", export.ZOOMS)
def test_svg_structure(store: Store, facts: list[Fact], zoom: str) -> None:
    svg = export.to_svg(store, CASE, facts, zoom=zoom)
    root = ET.fromstring(svg)  # well-formed despite <, &, " in titles
    assert root.tag == f"{SVG_NS}svg"
    assert export.MIN_WIDTH <= int(root.get("width")) <= export.MAX_WIDTH
    assert "<Acme> &" not in svg and "&lt;Acme&gt; &amp;" in svg
    # self-contained: no external fonts, stylesheets, images or links
    assert not [el for el in root.iter() if any("href" in a for a in el.attrib)]
    assert not [el for el in root.iter() if el.tag in (f"{SVG_NS}image", f"{SVG_NS}style")]

    facts_g = {g.get("data-fact-id"): g for g in _groups(root, "fact")}
    assert set(facts_g) == {f.fact_id for f in facts}
    key, rng, later, rel, undated = facts
    # key fact highlighted, disputed marked
    assert "key-fact" in facts_g[key.fact_id].get("class")
    assert "disputed" in facts_g[key.fact_id].get("class")
    marker = facts_g[key.fact_id].find(f"{SVG_NS}circle")
    assert marker.get("fill") == export.COLOR_KEY_FILL
    assert marker.get("stroke-dasharray")
    assert "(disputed)" in _texts(facts_g[key.fact_id])[0]
    # points for single dates, bars for ranges
    assert "range" in facts_g[rng.fact_id].get("class")
    assert facts_g[rng.fact_id].find(f"{SVG_NS}rect") is not None
    assert "[ALLEGED]" in _texts(facts_g[rng.fact_id])[0]
    assert "point" in facts_g[later.fact_id].get("class")
    # undated lane holds the facts without a resolved date
    lane = _groups(root, "undated-lane")
    assert len(lane) == 1
    assert any(t.startswith("Undated") for t in _texts(lane[0]))
    lane_ids = {g.get("data-fact-id") for g in _groups(lane[0], "fact")}
    assert lane_ids == {rel.fact_id, undated.fact_id}
    # ticks on the axis
    ticks = [t for t in root.iter(f"{SVG_NS}text") if t.get("class") == "tick"]
    assert ticks


def test_svg_ticks_match_zoom(store: Store, facts: list[Fact]) -> None:
    def ticks(zoom: str) -> list[str]:
        root = ET.fromstring(export.to_svg(store, CASE, facts, zoom=zoom))
        return [t.text for t in root.iter(f"{SVG_NS}text") if t.get("class") == "tick"]

    assert ticks("year") == ["2022"]
    assert ticks("month")[0] == "Mar 2022" and "Sep 2022" in ticks("month")
    assert ticks("week")[0] == "Feb 28, 2022"  # weeks start on Monday
    assert ticks("day")[0] == "Mar 1"


def test_svg_labels_do_not_overlap(store: Store) -> None:
    many = [
        store.add_fact(
            CASE,
            f"Event number {i} with a fairly long title",
            by="alice",
            date_start=date(2023, 1, 1 + i % 3),
        )
        for i in range(12)
    ]
    root = ET.fromstring(export.to_svg(store, CASE, many, zoom="month"))
    ys = [float(g.find(f"{SVG_NS}circle").get("cy")) for g in _groups(root, "point")]
    assert len(set(ys)) == 12  # same-day facts stack into separate rows


def test_svg_width_is_bounded(store: Store) -> None:
    wide = [
        store.add_fact(CASE, "Start", by="alice", date_start=date(1990, 1, 1)),
        store.add_fact(CASE, "End", by="alice", date_start=date(2024, 1, 1)),
    ]
    root = ET.fromstring(export.to_svg(store, CASE, wide, zoom="day"))
    assert int(root.get("width")) <= export.MAX_WIDTH
    one = [store.add_fact(CASE, "Only", by="alice", date_start=date(2024, 1, 1))]
    root = ET.fromstring(export.to_svg(store, CASE, one, zoom="year"))
    assert int(root.get("width")) >= export.MIN_WIDTH


def test_svg_only_undated(store: Store) -> None:
    only = [store.add_fact(CASE, "No date here", by="alice")]
    root = ET.fromstring(export.to_svg(store, CASE, only))
    assert len(_groups(root, "undated-lane")) == 1
    assert "No dated facts" in "".join(_texts(root))


def test_svg_strips_control_characters(store: Store) -> None:
    bad = [store.add_fact(CASE, "Bell \x07 and nul-ish \x1f text", by="alice")]
    ET.fromstring(export.to_svg(store, CASE, bad))


def test_svg_empty_and_bad_zoom(store: Store, facts: list[Fact]) -> None:
    root = ET.fromstring(export.to_svg(store, CASE, []))
    assert "No facts" in _texts(root)
    with pytest.raises(ValueError):
        export.to_svg(store, CASE, facts, zoom="decade")


def test_png(store: Store, facts: list[Fact]) -> None:
    png = export.to_png(store, CASE, facts, zoom="month")
    assert png.startswith(b"\x89PNG\r\n\x1a\n")
    empty = export.to_png(store, CASE, [])
    assert empty.startswith(b"\x89PNG\r\n\x1a\n")
