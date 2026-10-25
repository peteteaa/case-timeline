"""Discover evidence files and turn them into normalized text with page offsets.

Evidence files are only ever opened for reading. All offsets index into the
normalized text (NFC, ``\\n`` line endings, no NUL characters).
"""

from __future__ import annotations

import codecs
import email
import email.policy
import hashlib
import io
import logging
import os
import re
import unicodedata
from datetime import date
from email.message import EmailMessage, Message
from email.utils import parsedate_to_datetime
from html.parser import HTMLParser
from pathlib import Path

from pipeline.contracts import Document, Page, ParsedDocument

log = logging.getLogger(__name__)

MIME_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".eml": "message/rfc822",
}
SUPPORTED_EXTENSIONS = frozenset(MIME_TYPES)
PAGE_SEPARATOR = "\n\n"
_DOC_DATE_SCAN_LINES = 25
_HEADER_FIELDS = ("Date", "From", "To", "Cc", "Subject")

# --------------------------------------------------------------------------- normalization


def normalize_text(text: str) -> str:
    """NFC-normalize, unify line endings to ``\\n`` and drop NUL characters."""
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return unicodedata.normalize("NFC", text)


def _hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _hash_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------------------- discovery


def _default_inbox(case_id: str) -> list[tuple[Path, bool]]:
    """Settings inbox as (root, recursive): loose ``inbox/`` files plus all of ``inbox/<case>``."""
    try:
        from pipeline.config import get_settings, valid_case_id

        valid_case_id(case_id)
        inbox = Path(get_settings().inbox)
    except Exception:  # settings unavailable or misconfigured: no inbox
        log.debug("no inbox from settings", exc_info=True)
        return []
    # The bare inbox is scanned flat so other cases' subfolders are not picked up.
    return [(inbox, False), (inbox / case_id, True)]


def _candidate_files(root: Path, recursive: bool) -> list[Path]:
    if not root.is_dir():
        return []
    entries = root.rglob("*") if recursive else root.iterdir()
    out = []
    for p in entries:
        rel_parts = p.relative_to(root).parts
        if any(part.startswith((".", "~$")) for part in rel_parts):
            continue  # hidden files, Office lock files
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS:
            out.append(p)
    return sorted(out)


def discover(
    evidence_dir: Path, case_id: str, *, inbox_dirs: list[Path] | None = None
) -> list[Document]:
    """Hash every supported file under ``evidence_dir`` and the upload inbox.

    ``inbox_dirs`` (scanned recursively) defaults to the settings inbox: loose files in
    ``inbox/`` plus everything in ``inbox/<case_id>/``. Identical files are listed once.
    """
    evidence_dir = Path(evidence_dir)
    if not evidence_dir.is_dir():
        log.warning("evidence dir %s does not exist", evidence_dir)
    roots: list[tuple[Path, bool]] = [(evidence_dir, True)]
    if inbox_dirs is None:
        roots += _default_inbox(case_id)
    else:
        roots += [(Path(d), True) for d in inbox_dirs]

    docs: list[Document] = []
    seen: set[str] = set()
    for root, recursive in roots:
        for path in _candidate_files(root, recursive):
            try:
                sha = _hash_file(path)
            except OSError:
                log.warning("cannot read %s; skipping", path, exc_info=True)
                continue
            if sha in seen:
                continue
            seen.add(sha)
            docs.append(
                Document(
                    id=sha[:16],
                    case_id=case_id,
                    path=str(path.resolve()),
                    filename=path.name,
                    mime=MIME_TYPES[path.suffix.lower()],
                    sha256=sha,
                )
            )
    return docs


# --------------------------------------------------------------------------- parsing


def _join_pages(page_texts: list[str]) -> tuple[str, list[Page]]:
    """Normalize each page, join with PAGE_SEPARATOR and record offsets."""
    parts: list[str] = []
    pages: list[Page] = []
    pos = 0
    for i, raw in enumerate(page_texts, start=1):
        txt = normalize_text(raw).strip("\n")
        if parts:
            parts.append(PAGE_SEPARATOR)
            pos += len(PAGE_SEPARATOR)
        pages.append(Page(number=i, char_start=pos, char_end=pos + len(txt)))
        parts.append(txt)
        pos += len(txt)
    return "".join(parts), pages


def _docling_enabled() -> bool:
    return os.environ.get("CT_DISABLE_DOCLING", "").lower() not in {"1", "true", "yes"}


def _parse_pdf_docling(path: Path) -> list[str] | None:
    """Per-page markdown via Docling, or None if Docling is unavailable or fails."""
    if not _docling_enabled():
        return None
    try:
        from docling.document_converter import DocumentConverter  # type: ignore[import-not-found]
    except ImportError:
        return None
    try:
        result = DocumentConverter().convert(str(path))
        ddoc = result.document
        n = len(ddoc.pages)
        pages = [ddoc.export_to_markdown(page_no=i) for i in range(1, n + 1)]
    except Exception:
        log.warning("docling failed on %s; falling back to PyMuPDF", path.name, exc_info=True)
        return None
    return pages if any(p.strip() for p in pages) else None


def _parse_pdf_pymupdf(data: bytes) -> list[str]:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        return [page.get_text("text") for page in pdf]


def _parse_docx(data: bytes) -> str:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                cells = [c.text.strip() for c in row.cells]
                blocks.append(" | ".join(cells))
        else:
            blocks.append(item.text)
    return "\n".join(blocks)


class _HTMLText(HTMLParser):
    _BLOCK = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "blockquote"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    return re.sub(r"\n{3,}", "\n\n", "".join(p.parts)).strip()


def _email_text(msg: Message) -> str:
    """Headers plus body text; nested message/rfc822 parts are rendered recursively."""
    lines = [f"{h}: {msg[h]}" for h in _HEADER_FIELDS if msg[h]]
    body_parts: list[str] = []
    attachments: list[str] = []
    html_fallback: list[str] = []
    parts = msg.walk() if msg.is_multipart() else [msg]
    for part in parts:
        ctype = part.get_content_type()
        if part.is_multipart():
            if ctype == "message/rfc822":
                for inner in part.get_payload():
                    body_parts.append("\n" + _email_text(inner))
            continue
        filename = part.get_filename()
        if filename or part.get_content_disposition() == "attachment":
            attachments.append(filename or "(unnamed)")
            continue
        if ctype in {"text/plain", "text/html"}:
            try:
                content = part.get_content() if isinstance(part, EmailMessage) else None
            except (LookupError, ValueError):
                content = None
            if content is None:
                payload = part.get_payload(decode=True) or b""
                try:
                    content = payload.decode(part.get_content_charset() or "utf-8", "replace")
                except LookupError:  # bogus/unknown charset label: best-effort decode
                    content = _decode_text(payload)
            (body_parts if ctype == "text/plain" else html_fallback).append(content)
    if not body_parts:
        body_parts = [_html_to_text(h) for h in html_fallback]
    text = "\n".join(lines) + "\n\n" + "\n".join(b.strip("\n") for b in body_parts)
    if attachments:
        text += "\n\n" + "\n".join(f"[Attachment: {a}]" for a in attachments)
    return text


def _parse_eml(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    return _email_text(msg)


def _decode_text(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return data.decode("utf-32", "replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", "replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def parse(doc: Document) -> ParsedDocument:
    """Extract normalized text (and PDF page offsets) from a discovered document."""
    path = Path(doc.path)
    data = path.read_bytes()  # read-only access to evidence
    if _hash_bytes(data) != doc.sha256:
        log.warning("%s changed since discovery (sha256 mismatch)", doc.filename)
    ext = path.suffix.lower()
    pages: list[Page] = []
    if not data:  # zero-byte file: nothing to parse (and not a corrupt PDF/DOCX)
        text = ""
    elif ext == ".pdf":
        page_texts = _parse_pdf_docling(path) or _parse_pdf_pymupdf(data)
        text, pages = _join_pages(page_texts)
        doc = doc.model_copy(update={"num_pages": len(pages)})
    elif ext == ".docx":
        text = normalize_text(_parse_docx(data))
    elif ext == ".eml":
        text = normalize_text(_parse_eml(data))
    elif ext in {".txt", ".md"}:
        text = normalize_text(_decode_text(data))
    else:
        raise ValueError(f"unsupported file type: {doc.filename}")
    if not text.strip():
        log.warning("%s produced no text (scanned PDF without OCR?)", doc.filename)
    return ParsedDocument(doc=doc, text=text, pages=pages)


# --------------------------------------------------------------------------- document date

_MONTHS = (
    r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|"
    r"Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\.?"
)
_DATE_RE = re.compile(
    rf"\b(?:{_MONTHS}\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}}"
    rf"|\d{{1,2}}(?:st|nd|rd|th)?\s+{_MONTHS},?\s+\d{{4}}"
    r"|\d{4}-\d{2}-\d{2}"
    r"|\d{1,2}/\d{1,2}/\d{4})\b",
    re.IGNORECASE,
)
_DATE_LINE_RE = re.compile(r"^\W*(?:date|dated|sent)\b\s*[:\-]?\s*(.+)$", re.IGNORECASE)


def _parse_date_str(s: str) -> date | None:
    import dateparser

    dt = dateparser.parse(
        s,
        languages=["en"],
        settings={"PREFER_DAY_OF_MONTH": "first", "STRICT_PARSING": True},
    )
    return dt.date() if dt else None


def _email_header_date(value: str) -> date | None:
    try:
        return parsedate_to_datetime(value.strip()).date()
    except (TypeError, ValueError, IndexError):
        return _parse_date_str(value)


def doc_date(parsed: ParsedDocument) -> date | None:
    """Best-guess authoring date: EML ``Date`` header, else a date-like line near the top."""
    lines = parsed.text.split("\n")[:_DOC_DATE_SCAN_LINES]
    if parsed.doc.filename.lower().endswith(".eml"):
        for line in lines:
            if line.startswith("Date:"):
                if found := _email_header_date(line[5:]):
                    return found
                break
    # Prefer explicit "Date: ..." lines, then any line with a full calendar date.
    for line in lines:
        m = _DATE_LINE_RE.match(line)
        if m and (dm := _DATE_RE.search(m.group(1))):
            if found := _parse_date_str(dm.group(0)):
                return found
    for line in lines:
        if dm := _DATE_RE.search(line):
            if found := _parse_date_str(dm.group(0)):
                return found
    return None
