"""Exports (SPEC §7): XLSX/CSV of the current list, SVG/PNG of the visual timeline.

Every function takes the facts already filtered/sorted by the caller (the view's query), so
exports respect it. Spreadsheet columns come in the caller's on-screen order. vis-timeline
renders HTML, so the image exports are drawn here from the facts (SPEC §9.11): a time axis at
the chosen zoom, points for single dates, bars for ranges, key facts highlighted, disputed
facts marked, and undated facts in their own lane at the bottom.
"""

from __future__ import annotations

import csv
import io
import math
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, timedelta
from xml.sax.saxutils import escape

from pipeline.schema import Fact
from pipeline.store import Store

# Column key -> header, in the list view's default on-screen order.
COLUMNS: dict[str, str] = {
    "date": "Date",
    "title": "Title",
    "statuses": "Statuses",
    "review": "Review",
    "evidence": "Evidence",
    "tags": "Tags",
    "witnesses": "Witnesses",
    "notes": "Notes",
}
DEFAULT_COLUMNS: tuple[str, ...] = tuple(COLUMNS)
ZOOMS = ("day", "week", "month", "year")

STATUS_LABELS = {"key_fact": "Key fact", "case_impact": "Case impact", "disputed": "Disputed"}
REVIEW_LABELS = {"draft": "Draft", "accepted": "Accepted", "rejected": "Rejected"}
ALLEGED_PREFIX = "[ALLEGED] "

# --------------------------------------------------------------------------- cells


def format_date(fact: Fact) -> str:
    """The list view's date cell: start, "start – end", "(approx.)", date_text, or Undated."""
    if fact.date_start is None:
        return (fact.date_text or "").strip() or "Undated"
    text = fact.date_start.isoformat()
    if fact.date_end is not None and fact.date_end != fact.date_start:
        text += f" – {fact.date_end.isoformat()}"
    if fact.date_certainty == "approximate":
        text += " (approx.)"
    return text


def _check_columns(columns: Sequence[str]) -> list[str]:
    cols = list(columns)
    if not cols:
        raise ValueError("choose at least one column to export")
    unknown = [c for c in cols if c not in COLUMNS]
    if unknown:
        raise ValueError(f"unknown export columns: {unknown}; choose from {list(COLUMNS)}")
    return cols


@dataclass
class _Context:
    alleged: set[str]
    counts: dict[str, dict[str, int]]
    tags: dict[str, str]
    witnesses: dict[str, str]


def _context(store: Store, case_id: str, facts: Sequence[Fact]) -> _Context:
    return _Context(
        alleged=store.alleged_fact_ids(case_id),
        counts=store.fact_counts([f.fact_id for f in facts]),
        tags={t.tag_id: t.name for t in store.list_tags(case_id)},
        witnesses={w.witness_id: w.name for w in store.list_witnesses(case_id)},
    )


def _title(fact: Fact, ctx: _Context) -> str:
    return (ALLEGED_PREFIX if fact.fact_id in ctx.alleged else "") + fact.title


def _cell(key: str, fact: Fact, ctx: _Context) -> str | int:
    counts = ctx.counts.get(fact.fact_id, {})
    if key == "date":
        return format_date(fact)
    if key == "title":
        return _title(fact, ctx)
    if key == "statuses":
        return "; ".join(STATUS_LABELS.get(s, s) for s in fact.statuses)
    if key == "review":
        return REVIEW_LABELS.get(fact.review, fact.review)
    if key == "evidence":
        return counts.get("evidence", len(fact.evidence))
    if key == "tags":
        return "; ".join(ctx.tags.get(t, t) for t in fact.tag_ids)
    if key == "witnesses":
        return "; ".join(ctx.witnesses.get(w, w) for w in fact.witness_ids)
    if key == "notes":
        return counts.get("notes", 0)
    raise ValueError(f"unknown export column {key!r}")  # pragma: no cover - checked earlier


def _rows(
    store: Store, case_id: str, facts: Sequence[Fact], columns: Sequence[str]
) -> tuple[list[str], list[list[str | int]]]:
    cols = _check_columns(columns)
    ctx = _context(store, case_id, facts)
    return cols, [[_cell(c, f, ctx) for c in cols] for f in facts]


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(value: str | int) -> str | int:
    """Neutralize spreadsheet formulas in text that came from documents (CSV injection)."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


# --------------------------------------------------------------------------- CSV / XLSX


def to_csv(
    store: Store, case_id: str, facts: Sequence[Fact], columns: Sequence[str] = DEFAULT_COLUMNS
) -> bytes:
    """UTF-8 CSV with a BOM (so Excel reads it as UTF-8), header row first."""
    cols, rows = _rows(store, case_id, facts, columns)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([COLUMNS[c] for c in cols])
    for row in rows:
        writer.writerow([_csv_safe(v) for v in row])
    return buf.getvalue().encode("utf-8-sig")


_WIDTHS = {  # default / maximum character widths per column
    "date": (14, 34),
    "title": (30, 80),
    "statuses": (10, 34),
    "review": (9, 12),
    "evidence": (9, 10),
    "tags": (10, 40),
    "witnesses": (12, 40),
    "notes": (7, 8),
}
KEY_FACT_FILL = "FFF3C4"
_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def to_xlsx(
    store: Store, case_id: str, facts: Sequence[Fact], columns: Sequence[str] = DEFAULT_COLUMNS
) -> bytes:
    """One "Facts" sheet: bold frozen header, filters, key-fact rows on a light fill."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols, rows = _rows(store, case_id, facts, columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "Facts"
    bold = Font(bold=True)
    for j, key in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=COLUMNS[key])
        cell.font = bold
    fill = PatternFill(fill_type="solid", start_color=KEY_FACT_FILL, end_color=KEY_FACT_FILL)
    top = Alignment(vertical="top", wrap_text=False)
    wrap = Alignment(vertical="top", wrap_text=True)
    for i, (fact, row) in enumerate(zip(facts, rows, strict=True), start=2):
        key = "key_fact" in fact.statuses
        for j, (col, value) in enumerate(zip(cols, row, strict=True), start=1):
            cell = ws.cell(row=i, column=j)
            if isinstance(value, str):
                cell.value = _ILLEGAL_XLSX.sub("", value)
                cell.data_type = "s"  # never a formula, whatever the text starts with
            else:
                cell.value = value
            cell.alignment = wrap if col == "title" else top
            if key:
                cell.fill = fill
    for j, col in enumerate(cols, start=1):
        lo, hi = _WIDTHS[col]
        longest = max([len(COLUMNS[col])] + [len(str(r[j - 1])) for r in rows])
        ws.column_dimensions[get_column_letter(j)].width = max(lo, min(hi, longest + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows) + 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- SVG / PNG

FONT = "DejaVu Sans, Helvetica, Arial, sans-serif"  # system font names only; nothing fetched
MIN_WIDTH = 640
MAX_WIDTH = 8000
MARGIN = 24
FONT_PX = 12
CHAR_PX = 7.0  # rough average advance at FONT_PX, used for label layout
ROW_H = 24
LABEL_CHARS = 48

COLOR_TEXT = "#1f2933"
COLOR_MUTED = "#616e7c"
COLOR_AXIS = "#9aa5b1"
COLOR_GRID = "#e4e7eb"
COLOR_ITEM_FILL = "#d9e6f6"
COLOR_ITEM_STROKE = "#3d6eb0"
COLOR_KEY_FILL = "#ffd86b"
COLOR_KEY_STROKE = "#c48a00"
COLOR_DISPUTED = "#c62828"
COLOR_BG = "#ffffff"
COLOR_UNDATED_BG = "#f5f7fa"

# zoom -> pixels per tick unit (day, week, month, year) before fitting into the width bounds
_ZOOM_SCALE = {"day": 60.0, "week": 84.0, "month": 96.0, "year": 96.0}
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _x(text: str) -> str:
    """XML-escape text for element content and attribute values."""
    return escape(_INVALID_XML.sub("", text), {'"': "&quot;", "'": "&apos;"})


def _truncate(text: str, n: int = LABEL_CHARS) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _text_px(text: str) -> float:
    return len(text) * CHAR_PX


def _floor_unit(d: date, zoom: str) -> date:
    if zoom == "day":
        return d
    if zoom == "week":
        return d - timedelta(days=d.weekday())
    if zoom == "month":
        return d.replace(day=1)
    return d.replace(month=1, day=1)


def _next_unit(d: date, zoom: str) -> date:
    if zoom == "day":
        return d + timedelta(days=1)
    if zoom == "week":
        return d + timedelta(days=7)
    if zoom == "month":
        return date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return date(d.year + 1, 1, 1)


def _tick_label(d: date, zoom: str) -> str:
    if zoom == "year":
        return str(d.year)
    if zoom == "month":
        return d.strftime("%b %Y")
    label = f"{d.strftime('%b')} {d.day}"
    return f"{label}, {d.year}" if zoom == "week" else label


_UNIT_DAYS = {"day": 1.0, "week": 7.0, "month": 30.44, "year": 365.25}


@dataclass
class _Item:
    fact: Fact
    label: str
    key: bool
    disputed: bool
    x0: float = 0.0
    x1: float = 0.0  # bar end (== x0 for points)
    extent: float = 0.0  # rightmost pixel the item's label needs
    row: int = 0


def _label(fact: Fact, ctx: _Context) -> str:
    label = f"{fact.fact_id} · {_truncate(_title(fact, ctx))}"
    return label + (" (disputed)" if "disputed" in fact.statuses else "")


def _stack(items: list[_Item], gap: float = 10.0) -> int:
    """Greedy lane assignment by start x; returns the number of rows used."""
    row_ends: list[float] = []
    for item in sorted(items, key=lambda i: (i.x0, i.extent)):
        for r, end in enumerate(row_ends):
            if end + gap <= item.x0:
                item.row, row_ends[r] = r, item.extent
                break
        else:
            item.row = len(row_ends)
            row_ends.append(item.extent)
    return len(row_ends)


def _marker_colors(item: _Item) -> tuple[str, str, str]:
    """(fill, stroke, dash) for an item."""
    fill = COLOR_KEY_FILL if item.key else COLOR_ITEM_FILL
    stroke = COLOR_KEY_STROKE if item.key else COLOR_ITEM_STROKE
    if item.disputed:
        return fill, COLOR_DISPUTED, "4 2"
    return fill, stroke, ""


def _classes(item: _Item, *extra: str) -> str:
    names = ["fact", *extra]
    if item.key:
        names.append("key-fact")
    if item.disputed:
        names.append("disputed")
    return " ".join(names)


def _empty_svg(message: str = "No facts") -> str:
    w, h = 320, 96
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" '
        f'viewBox="0 0 {w} {h}" font-family="{FONT}">'
        f'<rect width="{w}" height="{h}" fill="{COLOR_BG}"/>'
        f'<text x="{w / 2}" y="{h / 2 + 5}" text-anchor="middle" font-size="16" '
        f'fill="{COLOR_MUTED}">{_x(message)}</text></svg>'
    )


def to_svg(store: Store, case_id: str, facts: Sequence[Fact], zoom: str = "month") -> str:
    """A self-contained SVG of the visual timeline at ``zoom`` (day, week, month or year)."""
    if zoom not in ZOOMS:
        raise ValueError(f"unknown zoom {zoom!r}; choose from {list(ZOOMS)}")
    if not facts:
        return _empty_svg()
    ctx = _context(store, case_id, facts)
    dated = [f for f in facts if f.date_start is not None]
    undated = [f for f in facts if f.date_start is None]

    def make(f: Fact) -> _Item:
        return _Item(f, _label(f, ctx), "key_fact" in f.statuses, "disputed" in f.statuses)

    # ---- horizontal scale
    items = [make(f) for f in dated]
    plot_x = MARGIN
    if dated:
        lo = _floor_unit(min(f.date_start for f in dated), zoom)  # type: ignore[type-var]
        last = max((f.date_end or f.date_start) for f in dated)  # type: ignore[type-var]
        hi = _next_unit(_floor_unit(last, zoom), zoom)
        span_days = max((hi - lo).days, 1)
        px_per_day = _ZOOM_SCALE[zoom] / _UNIT_DAYS[zoom]
        plot_w = span_days * px_per_day
        if plot_w > MAX_WIDTH - 2 * MARGIN:
            plot_w = MAX_WIDTH - 2 * MARGIN
            px_per_day = plot_w / span_days
        if plot_w < MIN_WIDTH - 2 * MARGIN:
            plot_w = MIN_WIDTH - 2 * MARGIN
            px_per_day = plot_w / span_days

        def xpos(d: date) -> float:
            return plot_x + (d - lo).days * px_per_day

        for it in items:
            f = it.fact
            it.x0 = xpos(f.date_start)  # type: ignore[arg-type]
            if f.date_end is not None and f.date_end > f.date_start:  # type: ignore[operator]
                it.x1 = max(xpos(f.date_end + timedelta(days=1)), it.x0 + 8)
                it.extent = max(it.x1, it.x0 + 6 + _text_px(it.label))
            else:
                it.x1 = it.x0
                it.extent = it.x0 + 9 + _text_px(it.label)
        rows = _stack(items)
        axis_end = plot_x + plot_w
    else:
        lo = hi = None
        px_per_day = 0.0
        rows = 0
        axis_end = MIN_WIDTH - MARGIN

    # ---- undated lane: wrap chips left to right
    content_right = max([axis_end, *(i.extent for i in items)])
    width = int(min(MAX_WIDTH, max(MIN_WIDTH, math.ceil(content_right + MARGIN))))
    chips: list[tuple[_Item, float, int, float]] = []  # item, x, row, chip width
    cx, crow = float(MARGIN), 0
    for f in undated:
        it = make(f)
        cw = 22 + _text_px(it.label)
        if cx + cw > width - MARGIN and cx > MARGIN:
            cx, crow = float(MARGIN), crow + 1
        chips.append((it, cx, crow, cw))
        cx += cw + 10
    chip_rows = (crow + 1) if undated else 0

    # ---- vertical layout
    header_h = 34
    axis_y = header_h + 30
    rows_top = axis_y + 18
    dated_bottom = rows_top + rows * ROW_H + (8 if rows else 0)
    undated_top = dated_bottom + 16
    undated_h = (28 + chip_rows * ROW_H + 6) if undated else 0
    legend_y = undated_top + undated_h + 20
    height = legend_y + 18

    out: list[str] = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}" font-family="{FONT}" font-size="{FONT_PX}">',
        f'<rect width="{width}" height="{height}" fill="{COLOR_BG}"/>',
    ]
    n_dated, n_undated = len(dated), len(undated)
    summary = f"{len(facts)} fact{'s' * (len(facts) != 1)}"
    if dated:
        summary += f" · {lo.isoformat()} to {(hi - timedelta(days=1)).isoformat()}"  # type: ignore[union-attr,operator]
    summary += f" · zoom: {zoom}"
    out.append(
        f'<text x="{MARGIN}" y="22" font-size="15" font-weight="bold" fill="{COLOR_TEXT}">'
        f"{_x(f'Case {case_id} timeline')}</text>"
        f'<text x="{width - MARGIN}" y="22" text-anchor="end" fill="{COLOR_MUTED}">'
        f"{_x(summary)}</text>"
    )

    # ---- axis and ticks
    out.append('<g class="axis">')
    if dated:
        out.append(
            f'<line x1="{plot_x}" y1="{axis_y}" x2="{axis_end:.1f}" y2="{axis_y}" '
            f'stroke="{COLOR_AXIS}" stroke-width="1"/>'
        )
        ticks: list[date] = []
        t = lo
        while t <= hi:  # type: ignore[operator]
            ticks.append(t)  # type: ignore[arg-type]
            t = _next_unit(t, zoom)  # type: ignore[arg-type]
        tick_px = _UNIT_DAYS[zoom] * px_per_day
        label_px = max(_text_px(_tick_label(d, zoom)) for d in ticks) * 0.85 + 12
        step = max(1, math.ceil(label_px / max(tick_px, 0.01)))
        grid_bottom = max(dated_bottom, rows_top + 4)
        for k, d in enumerate(ticks):
            x = xpos(d)
            major = k % step == 0
            if major:
                out.append(
                    f'<line x1="{x:.1f}" y1="{axis_y}" x2="{x:.1f}" y2="{grid_bottom}" '
                    f'stroke="{COLOR_GRID}" stroke-width="1"/>'
                )
            if major or tick_px >= 3:
                out.append(
                    f'<line x1="{x:.1f}" y1="{axis_y - (6 if major else 3)}" x2="{x:.1f}" '
                    f'y2="{axis_y}" stroke="{COLOR_AXIS}" stroke-width="1"/>'
                )
            if major and d < hi:  # type: ignore[operator]
                out.append(
                    f'<text class="tick" x="{x + 3:.1f}" y="{axis_y - 9}" font-size="10" '
                    f'fill="{COLOR_MUTED}">{_x(_tick_label(d, zoom))}</text>'
                )
    else:
        out.append(
            f'<text x="{MARGIN}" y="{axis_y - 9}" fill="{COLOR_MUTED}">No dated facts</text>'
        )
    out.append("</g>")

    # ---- dated items
    out.append('<g class="dated">')
    for it in items:
        y = rows_top + it.row * ROW_H + ROW_H / 2
        fill, stroke, dash = _marker_colors(it)
        dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
        f = it.fact
        tooltip = f"{f.fact_id}: {_title(f, ctx)} ({format_date(f)})"
        if it.x1 > it.x0:
            out.append(
                f'<g class="{_classes(it, "range")}" data-fact-id="{_x(f.fact_id)}">'
                f"<title>{_x(tooltip)}</title>"
                f'<rect x="{it.x0:.1f}" y="{y - 8}" width="{it.x1 - it.x0:.1f}" height="16" '
                f'rx="3" fill="{fill}" stroke="{stroke}" stroke-width="1.5"{dash_attr}/>'
                f'<text x="{it.x0 + 5:.1f}" y="{y + 4}" fill="{COLOR_TEXT}">'
                f"{_x(it.label)}</text></g>"
            )
        else:
            out.append(
                f'<g class="{_classes(it, "point")}" data-fact-id="{_x(f.fact_id)}">'
                f"<title>{_x(tooltip)}</title>"
                f'<circle cx="{it.x0:.1f}" cy="{y}" r="5" fill="{fill}" stroke="{stroke}" '
                f'stroke-width="1.5"{dash_attr}/>'
                f'<text x="{it.x0 + 9:.1f}" y="{y + 4}" fill="{COLOR_TEXT}">'
                f"{_x(it.label)}</text></g>"
            )
    out.append("</g>")

    # ---- undated lane
    if undated:
        out.append(
            f'<g class="undated-lane"><rect x="{MARGIN / 2}" y="{undated_top}" '
            f'width="{width - MARGIN}" height="{undated_h}" rx="4" fill="{COLOR_UNDATED_BG}" '
            f'stroke="{COLOR_GRID}"/>'
            f'<text x="{MARGIN}" y="{undated_top + 18}" font-weight="bold" '
            f'fill="{COLOR_MUTED}">Undated ({n_undated})</text>'
        )
        for it, x, row, cw in chips:
            y = undated_top + 28 + row * ROW_H + ROW_H / 2
            fill, stroke, dash = _marker_colors(it)
            dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
            f = it.fact
            tooltip = f"{f.fact_id}: {_title(f, ctx)} ({format_date(f)})"
            out.append(
                f'<g class="{_classes(it, "undated")}" data-fact-id="{_x(f.fact_id)}">'
                f"<title>{_x(tooltip)}</title>"
                f'<rect x="{x:.1f}" y="{y - 9}" width="{cw:.1f}" height="18" rx="9" '
                f'fill="{fill}" stroke="{stroke}" stroke-width="1.5"{dash_attr}/>'
                f'<text x="{x + 11:.1f}" y="{y + 4}" fill="{COLOR_TEXT}">{_x(it.label)}</text>'
                "</g>"
            )
        out.append("</g>")

    # ---- legend
    lx = float(MARGIN)
    out.append(f'<g class="legend" font-size="11" fill="{COLOR_MUTED}">')
    for label, fill, stroke, dash in (
        ("Fact", COLOR_ITEM_FILL, COLOR_ITEM_STROKE, ""),
        ("Key fact", COLOR_KEY_FILL, COLOR_KEY_STROKE, ""),
        ("Disputed", COLOR_ITEM_FILL, COLOR_DISPUTED, "4 2"),
    ):
        dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
        out.append(
            f'<rect x="{lx}" y="{legend_y - 9}" width="18" height="10" rx="2" fill="{fill}" '
            f'stroke="{stroke}" stroke-width="1.5"{dash_attr}/>'
            f'<text x="{lx + 24}" y="{legend_y}">{label}</text>'
        )
        lx += 24 + _text_px(label) + 18
    out.append(
        f'<text x="{lx}" y="{legend_y}">● single date · ▬ date range · '
        f"{n_dated} dated, {n_undated} undated</text>"
    )
    out.append("</g></svg>")
    return "".join(out)


def to_png(store: Store, case_id: str, facts: Sequence[Fact], zoom: str = "month") -> bytes:
    """The SVG export rasterized with cairosvg."""
    import cairosvg

    svg = to_svg(store, case_id, facts, zoom)
    return cairosvg.svg2png(bytestring=svg.encode("utf-8"))
