"""Generate the FICTIONAL "demo" employment-dispute case and its gold labels.

All people, companies, addresses and events below are invented for testing.

    uv run python scripts/make_fixtures.py [--out tests/fixtures/case_demo] [--gold eval/gold]

Outputs are deterministic (fixed metadata and zip timestamps) so document ids stay stable.
This dev script is the only thing allowed to write the fixture folder; the pipeline never does.
"""

from __future__ import annotations

import argparse
import io
import json
import logging
import zipfile
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
log = logging.getLogger("make_fixtures")

PDF_NAME = "incident_report.pdf"
DOCX_NAME = "witness_statement_becker.docx"
EML_NAME = "email_thread.eml"
TXT_NAME = "deposition_excerpt_ellison.txt"
MD_NAME = "hr_memo.md"

RIVERA, ELLISON, NAND, BECKER = "Dana Rivera", "Mark Ellison", "Priya Nand", "Tom Becker"
ORTIZ, WHITFIELD = "Jamie Ortiz", "Karen Whitfield"

# --------------------------------------------------------------------------- document text

PDF_PAGES = [
    """LAKEVIEW POLICE DEPARTMENT
INCIDENT REPORT

Report No.: LPD-24-03318
Date of Report: March 18, 2024
Reporting Officer: Officer Jamie Ortiz, Badge 4471
Location: Halvorsen Freight Co., 2200 Dock Road, Lakeview (employee parking lot)
Incident Type: Verbal disturbance (no arrest)

NARRATIVE

On March 14, 2024 at approximately 5:40 PM, Officer Jamie Ortiz responded to the employee \
parking lot of Halvorsen Freight Co. regarding a verbal altercation between Dana Rivera and \
Mark Ellison.

Ms. Rivera stated that she had been sent home without pay earlier that day. She stated that \
on March 12, 2024, during a meeting in Mr. Ellison's office, Mr. Ellison threatened to fire \
her if she "kept writing up the forklifts."

Mr. Ellison stated that the March 12, 2024 meeting concerned Ms. Rivera's attendance and that \
only he and Ms. Rivera were present in the office. He denied making any threat.

No physical contact was reported. Both parties were advised and left the lot separately at \
approximately 6:15 PM.""",
    """LAKEVIEW POLICE DEPARTMENT - Report No. LPD-24-03318 (continued)

FOLLOW-UP

On March 15, 2024, Officer Ortiz conducted a follow-up telephone interview with Tom Becker, a \
forklift operator at Halvorsen Freight Co. Mr. Becker stated that he did not witness the \
parking lot incident but had heard raised voices from Mr. Ellison's office earlier in the week.

DISPOSITION

No criminal offense identified. Case closed; information only.

Report reviewed and approved by Sgt. Karen Whitfield on March 18, 2024.""",
]

DOCX_TITLE = "WITNESS STATEMENT OF TOM BECKER"
DOCX_PARAGRAPHS = [
    "Matter: Rivera v. Halvorsen Freight Co.",
    "Date: April 2, 2024",
    "1. My name is Tom Becker. I have worked as a forklift operator at Halvorsen Freight Co. "
    "since 2019.",
    "2. In early March 2024, Dana Rivera told me that the brakes on forklift #7 were failing "
    "and that she had reported it to Mark Ellison.",
    "3. The meeting in Mark Ellison's office took place on March 11, 2024. I was in the room "
    "for the whole meeting because Dana asked me to come as a witness. Mark told Dana that if "
    "she kept filing safety reports she would be “done here.”",
    "4. Two days after the meeting, Dana was sent home without pay.",
    "5. Forklift #7 was taken out of service on March 22, 2024 after its brakes failed during "
    "a shift.",
    "6. I make this statement freely and believe it to be true.",
]
DOCX_SIGNATURE = [("Signed", "Tom Becker"), ("Date signed", "April 2, 2024")]

EML_MESSAGES = [
    # (from, to, sent header, subject, body) — oldest last, as in a reply chain
    (
        "Priya Nand <pnand@halvorsen-freight.example>",
        "Dana Rivera <dana.rivera@mail.example>",
        "Mon, 25 Mar 2024 09:14:00 -0400",
        "RE: Formal complaint - safety retaliation",
        """Dana,

Due to a scheduling conflict, your investigation interview has been moved from March 27, 2024 \
to March 28, 2024 at 10:00 AM in Conference Room B. Please let me know if you would like a \
representative present.

Priya Nand
HR Manager, Halvorsen Freight Co.""",
    ),
    (
        "Priya Nand <pnand@halvorsen-freight.example>",
        "Dana Rivera <dana.rivera@mail.example>",
        "Wednesday, March 20, 2024 4:02 PM",
        "RE: Formal complaint - safety retaliation",
        """Dana,

I confirm receipt of your complaint dated March 18, 2024. HR has opened an investigation and \
scheduled an interview with you for March 27, 2024.

Priya""",
    ),
    (
        "Dana Rivera <dana.rivera@mail.example>",
        "Priya Nand <pnand@halvorsen-freight.example>",
        "Monday, March 18, 2024 7:48 AM",
        "Formal complaint - safety retaliation",
        """Priya,

I am filing a formal complaint. On March 4, 2024 I reported in writing to Mark Ellison that \
the brakes on forklift #7 were failing. On March 12, 2024 he called me into his office and \
told me to stop filing safety reports. Two days later I was sent home without pay.

I am asking HR to investigate.

Dana Rivera""",
    ),
]

TXT_TEXT = """DEPOSITION OF MARK ELLISON (EXCERPT)
Rivera v. Halvorsen Freight Co.
Taken September 10, 2024, at the offices of Brandt & Okafor LLP, Lakeview

[Page 14]

Q. When did your meeting with Ms. Rivera take place?
A. March 12, 2024. In the afternoon, I believe.
Q. Who else was present at that meeting?
A. Priya Nand from HR sat in. Nobody else.
Q. Was Tom Becker in the room?
A. No.

[Page 15]

Q. When did you first learn about problems with the brakes on forklift #7?
A. Sometime in mid-March, after that meeting. Somebody mentioned it on the floor.
Q. You did not receive a written report from Ms. Rivera before that meeting?
A. I don't recall receiving one.
Q. When was Ms. Rivera's employment terminated?
A. April 5, 2024.
Q. Who made that decision?
A. I did, together with HR.
"""

MD_TEXT = """# MEMORANDUM TO FILE

**From:** Priya Nand, HR Manager
**Date:** April 8, 2024
**Re:** Dana Rivera — attendance, complaint investigation, and separation

## Summary of events

- On March 12, 2024, I attended a meeting between Mark Ellison and Dana Rivera regarding \
attendance concerns. No one else was present.
- Ms. Rivera was placed on unpaid administrative leave on March 14, 2024 pending review.
- Ms. Rivera submitted a formal complaint alleging safety retaliation on March 18, 2024.
- The investigation interview with Ms. Rivera was held on March 28, 2024.
- Ms. Rivera's employment was terminated effective April 5, 2024.

## Appeal

Under Policy 4.2, Ms. Rivera may appeal the termination in writing. The appeal must be \
received by May 6, 2024.

## Addendum — June 3, 2024

On June 3, 2024, the company received notice that Ms. Rivera filed a charge of retaliation \
with the State Civil Rights Division on May 28, 2024. The company's position statement is due \
by July 1, 2024. Legal counsel has been notified.
"""

# --------------------------------------------------------------------------- gold labels


def _g(
    key: str,
    doc: str,
    date: str | None,
    date_text: str,
    certainty: str,
    actors: list[str],
    description: str,
    quote: str,
    *,
    date_range: tuple[str, str] | None = None,
    deadline: bool = False,
) -> dict:
    item = {
        "date": date,
        "actors": actors,
        "description": description,
        "doc": doc,
        "event_key": key,
        "date_text": date_text,
        "certainty": certainty,
        "deadline_relevant": deadline,
        "quote": quote,
    }
    if date_range:
        item["date_range"] = list(date_range)
    return item


GOLD = [
    # incident_report.pdf
    _g(
        "parking_lot_altercation",
        PDF_NAME,
        "2024-03-14",
        "March 14, 2024",
        "exact",
        [ORTIZ, RIVERA, ELLISON],
        "Police respond to a verbal altercation between Rivera and Ellison in the employer's "
        "parking lot.",
        "On March 14, 2024 at approximately 5:40 PM, Officer Jamie Ortiz responded",
    ),
    _g(
        "unpaid_leave",
        PDF_NAME,
        "2024-03-14",
        "earlier that day",
        "relative",
        [RIVERA],
        "Rivera says she was sent home without pay earlier that day.",
        "she had been sent home without pay earlier that day",
    ),
    _g(
        "office_meeting",
        PDF_NAME,
        "2024-03-12",
        "March 12, 2024",
        "exact",
        [RIVERA, ELLISON],
        "Meeting in Ellison's office; Rivera says Ellison threatened to fire her over forklift "
        "write-ups; Ellison says it concerned attendance and only the two of them were present.",
        "only he and Ms. Rivera were present in the office",
    ),
    _g(
        "police_followup",
        PDF_NAME,
        "2024-03-15",
        "March 15, 2024",
        "exact",
        [ORTIZ, BECKER],
        "Officer Ortiz conducts a follow-up phone interview with Becker, who heard raised voices "
        "from Ellison's office.",
        "On March 15, 2024, Officer Ortiz conducted a follow-up telephone interview",
    ),
    _g(
        "report_approved",
        PDF_NAME,
        "2024-03-18",
        "March 18, 2024",
        "exact",
        [WHITFIELD],
        "Police incident report LPD-24-03318 reviewed and approved; case closed.",
        "Report reviewed and approved by Sgt. Karen Whitfield on March 18, 2024.",
    ),
    # witness_statement_becker.docx
    _g(
        "brake_report_conversation",
        DOCX_NAME,
        None,
        "early March 2024",
        "approximate",
        [RIVERA, BECKER, ELLISON],
        "Rivera tells Becker forklift #7 brakes are failing and that she reported it to Ellison.",
        "In early March 2024, Dana Rivera told me that the brakes on forklift #7 were failing",
        date_range=("2024-03-01", "2024-03-10"),
    ),
    _g(
        "office_meeting",
        DOCX_NAME,
        "2024-03-11",
        "March 11, 2024",
        "exact",
        [RIVERA, ELLISON, BECKER],
        "Meeting in Ellison's office with Becker present; Ellison tells Rivera she will be "
        "'done here' if she keeps filing safety reports.",
        "The meeting in Mark Ellison's office took place on March 11, 2024.",
    ),
    _g(
        "unpaid_leave",
        DOCX_NAME,
        "2024-03-13",
        "Two days after the meeting",
        "relative",
        [RIVERA],
        "Rivera is sent home without pay two days after the meeting.",
        "Two days after the meeting, Dana was sent home without pay.",
    ),
    _g(
        "forklift_out_of_service",
        DOCX_NAME,
        "2024-03-22",
        "March 22, 2024",
        "exact",
        [],
        "Forklift #7 taken out of service after its brakes failed during a shift.",
        "Forklift #7 was taken out of service on March 22, 2024",
    ),
    _g(
        "becker_statement",
        DOCX_NAME,
        "2024-04-02",
        "April 2, 2024",
        "exact",
        [BECKER],
        "Becker signs his witness statement.",
        "Date signed | April 2, 2024",
    ),
    # email_thread.eml
    _g(
        "brake_report_written",
        EML_NAME,
        "2024-03-04",
        "March 4, 2024",
        "exact",
        [RIVERA, ELLISON],
        "Rivera reports in writing to Ellison that forklift #7 brakes are failing.",
        "On March 4, 2024 I reported in writing to Mark Ellison that the brakes on forklift #7 "
        "were failing.",
    ),
    _g(
        "office_meeting",
        EML_NAME,
        "2024-03-12",
        "March 12, 2024",
        "exact",
        [RIVERA, ELLISON],
        "Ellison calls Rivera into his office and tells her to stop filing safety reports.",
        "On March 12, 2024 he called me into his office and told me to stop filing safety reports.",
    ),
    _g(
        "unpaid_leave",
        EML_NAME,
        "2024-03-14",
        "Two days later",
        "relative",
        [RIVERA],
        "Rivera is sent home without pay two days after the meeting.",
        "Two days later I was sent home without pay.",
    ),
    _g(
        "hr_complaint",
        EML_NAME,
        "2024-03-18",
        "Monday, March 18, 2024",
        "exact",
        [RIVERA, NAND],
        "Rivera files a formal safety-retaliation complaint with HR.",
        "I am filing a formal complaint.",
    ),
    _g(
        "hr_ack",
        EML_NAME,
        "2024-03-20",
        "Wednesday, March 20, 2024",
        "exact",
        [NAND, RIVERA],
        "HR acknowledges the complaint, opens an investigation and schedules an interview for "
        "March 27, 2024.",
        "HR has opened an investigation and scheduled an interview with you for March 27, 2024.",
    ),
    _g(
        "investigation_interview",
        EML_NAME,
        "2024-03-28",
        "March 28, 2024",
        "exact",
        [NAND, RIVERA],
        "Investigation interview rescheduled (on March 25) from March 27 to March 28, 2024.",
        "your investigation interview has been moved from March 27, 2024 to March 28, 2024",
    ),
    # deposition_excerpt_ellison.txt
    _g(
        "ellison_deposition",
        TXT_NAME,
        "2024-09-10",
        "September 10, 2024",
        "exact",
        [ELLISON],
        "Deposition of Mark Ellison taken.",
        "Taken September 10, 2024",
    ),
    _g(
        "office_meeting",
        TXT_NAME,
        "2024-03-12",
        "March 12, 2024",
        "exact",
        [ELLISON, RIVERA, NAND],
        "Ellison testifies the meeting with Rivera was on March 12, 2024, attended by Nand; "
        "denies Becker was present.",
        "A. Priya Nand from HR sat in. Nobody else.",
    ),
    _g(
        "ellison_learns_brakes",
        TXT_NAME,
        None,
        "mid-March",
        "approximate",
        [ELLISON],
        "Ellison testifies he first learned of the forklift #7 brake problems in mid-March, "
        "after the meeting, and does not recall a written report.",
        "Sometime in mid-March, after that meeting.",
        date_range=("2024-03-11", "2024-03-20"),
    ),
    _g(
        "termination",
        TXT_NAME,
        "2024-04-05",
        "April 5, 2024",
        "exact",
        [ELLISON, RIVERA],
        "Rivera's employment terminated (decision by Ellison with HR).",
        "A. April 5, 2024.",
    ),
    # hr_memo.md
    _g(
        "office_meeting",
        MD_NAME,
        "2024-03-12",
        "March 12, 2024",
        "exact",
        [NAND, ELLISON, RIVERA],
        "Nand attends the meeting between Ellison and Rivera about attendance; no one else "
        "present.",
        "On March 12, 2024, I attended a meeting between Mark Ellison and Dana Rivera",
    ),
    _g(
        "unpaid_leave",
        MD_NAME,
        "2024-03-14",
        "March 14, 2024",
        "exact",
        [RIVERA],
        "Rivera placed on unpaid administrative leave pending review.",
        "Ms. Rivera was placed on unpaid administrative leave on March 14, 2024",
    ),
    _g(
        "hr_complaint",
        MD_NAME,
        "2024-03-18",
        "March 18, 2024",
        "exact",
        [RIVERA],
        "Rivera submits a formal complaint alleging safety retaliation.",
        "Ms. Rivera submitted a formal complaint alleging safety retaliation on March 18, 2024.",
    ),
    _g(
        "investigation_interview",
        MD_NAME,
        "2024-03-28",
        "March 28, 2024",
        "exact",
        [RIVERA, NAND],
        "Investigation interview with Rivera held.",
        "The investigation interview with Ms. Rivera was held on March 28, 2024.",
    ),
    _g(
        "termination",
        MD_NAME,
        "2024-04-05",
        "April 5, 2024",
        "exact",
        [RIVERA],
        "Rivera's employment terminated.",
        "Ms. Rivera's employment was terminated effective April 5, 2024.",
    ),
    _g(
        "appeal_deadline",
        MD_NAME,
        "2024-05-06",
        "May 6, 2024",
        "exact",
        [RIVERA],
        "Deadline for Rivera's written appeal of the termination under Policy 4.2.",
        "The appeal must be received by May 6, 2024.",
        deadline=True,
    ),
    _g(
        "agency_charge",
        MD_NAME,
        "2024-05-28",
        "May 28, 2024",
        "exact",
        [RIVERA],
        "Rivera files a retaliation charge with the State Civil Rights Division.",
        "filed a charge of retaliation with the State Civil Rights Division on May 28, 2024",
    ),
    _g(
        "charge_notice",
        MD_NAME,
        "2024-06-03",
        "June 3, 2024",
        "exact",
        [],
        "Company receives notice of Rivera's agency charge; legal counsel notified.",
        "On June 3, 2024, the company received notice",
    ),
    _g(
        "position_statement_due",
        MD_NAME,
        "2024-07-01",
        "July 1, 2024",
        "exact",
        [],
        "Company's position statement to the State Civil Rights Division is due.",
        "The company's position statement is due by July 1, 2024.",
        deadline=True,
    ),
]

CONTRADICTIONS = [
    {
        "id": "C1",
        "field": "date",
        "event_key": "office_meeting",
        "planted": True,
        "explanation": "Becker dates the office meeting March 11, 2024; the police report, "
        "Rivera's email, Ellison's deposition and the HR memo all say March 12, 2024.",
        "sides": [
            {"doc": DOCX_NAME, "quote": "took place on March 11, 2024"},
            {"doc": PDF_NAME, "quote": "on March 12, 2024, during a meeting in Mr. Ellison's"},
            {"doc": EML_NAME, "quote": "On March 12, 2024 he called me into his office"},
            {"doc": TXT_NAME, "quote": "A. March 12, 2024. In the afternoon, I believe."},
            {"doc": MD_NAME, "quote": "On March 12, 2024, I attended a meeting"},
        ],
    },
    {
        "id": "C2",
        "field": "actor",
        "event_key": "office_meeting",
        "planted": True,
        "explanation": "Who attended the meeting: Ellison told police only he and Rivera were "
        "present; Becker says he was in the room; Ellison's deposition and Nand's memo say Nand "
        "attended and no one else.",
        "sides": [
            {"doc": PDF_NAME, "quote": "only he and Ms. Rivera were present in the office"},
            {"doc": DOCX_NAME, "quote": "I was in the room for the whole meeting"},
            {"doc": TXT_NAME, "quote": "A. Priya Nand from HR sat in. Nobody else."},
            {"doc": MD_NAME, "quote": "I attended a meeting between Mark Ellison and Dana Rivera"},
        ],
    },
    {
        "id": "C3",
        "field": "fact",
        "event_key": "ellison_learns_brakes",
        "planted": True,
        "explanation": "Ellison testifies he first learned of the forklift #7 brake problem in "
        "mid-March, after the meeting; Rivera's email says she reported it to him in writing on "
        "March 4, 2024, and Becker says she had reported it by early March.",
        "sides": [
            {"doc": TXT_NAME, "quote": "Sometime in mid-March, after that meeting."},
            {"doc": EML_NAME, "quote": "On March 4, 2024 I reported in writing to Mark Ellison"},
            {"doc": DOCX_NAME, "quote": "she had reported it to Mark Ellison"},
        ],
    },
    {
        "id": "C4",
        "field": "date",
        "event_key": "unpaid_leave",
        "planted": False,
        "derived": True,
        "explanation": "Becker's 'two days after the meeting' resolves to March 13, 2024 using "
        "his own meeting date (March 11); the police report, Rivera's email and the HR memo put "
        "the unpaid leave on March 14, 2024. A consequence of C1.",
        "sides": [
            {
                "doc": DOCX_NAME,
                "quote": "Two days after the meeting, Dana was sent home without pay.",
            },
            {"doc": MD_NAME, "quote": "placed on unpaid administrative leave on March 14, 2024"},
        ],
    },
    {
        "id": "C5",
        "field": "fact",
        "event_key": "office_meeting",
        "planted": False,
        "derived": True,
        "explanation": "Becker told police on March 15 that he heard raised voices from "
        "Ellison's office (implying he was outside), but his statement says he was in the room "
        "for the whole meeting.",
        "sides": [
            {"doc": PDF_NAME, "quote": "had heard raised voices from Mr. Ellison's office"},
            {"doc": DOCX_NAME, "quote": "I was in the room for the whole meeting"},
        ],
    },
    {
        "id": "C6",
        "field": "fact",
        "event_key": "office_meeting",
        "planted": False,
        "derived": True,
        "explanation": "What the meeting was about: Rivera (email, and her account to police) "
        "and Becker say Ellison told her to stop filing safety reports; Ellison told police, and "
        "Nand's memo records, that it concerned attendance, and Ellison denied any threat.",
        "sides": [
            {"doc": EML_NAME, "quote": "told me to stop filing safety reports"},
            {"doc": DOCX_NAME, "quote": "if she kept filing safety reports she would be"},
            {
                "doc": PDF_NAME,
                "quote": "the March 12, 2024 meeting concerned Ms. Rivera's attendance",
            },
            {"doc": MD_NAME, "quote": "regarding attendance concerns"},
        ],
    },
]


def gold_timeline() -> list[dict]:
    """GOLD in the eval/run_eval.py fact layout (title, people, date_end, date_certainty)."""
    out = []
    for i, g in enumerate(GOLD, start=1):
        start, end = g["date_range"] if g.get("date_range") else (g["date"], None)
        out.append(
            {
                "id": f"G{i:02d}",
                "title": g["description"],
                "date": start,
                "date_end": end,
                "date_text": g["date_text"],
                "date_certainty": g["certainty"],
                "doc": g["doc"],
                "event_key": g["event_key"],
                "quote": g["quote"],
                "people": g["actors"],
                "deadline": g["deadline_relevant"],
            }
        )
    return out


def gold_disputes() -> list[dict]:
    """CONTRADICTIONS as gold dispute labels (every planted/derived one is disputed)."""
    return [{**c, "label": "disputed"} for c in CONTRADICTIONS]


# --------------------------------------------------------------------------- writers

FIXED_DT = datetime(2024, 10, 1, 12, 0, 0)


def make_pdf(path: Path) -> None:
    import pymupdf

    pdf = pymupdf.open()
    for text in PDF_PAGES:
        page = pdf.new_page(width=612, height=792)
        rect = pymupdf.Rect(72, 72, 540, 720)
        rc = page.insert_textbox(rect, text, fontname="helv", fontsize=11, lineheight=1.3)
        if rc < 0:
            raise RuntimeError(f"PDF page text overflow ({rc})")
    stamp = "D:20241001120000Z"
    pdf.set_metadata(
        {
            "title": "Lakeview PD Incident Report LPD-24-03318 (fictional)",
            "author": "make_fixtures.py",
            "creationDate": stamp,
            "modDate": stamp,
            "producer": "",
            "creator": "",
        }
    )
    pdf.save(path, garbage=4, deflate=True, no_new_id=True)
    pdf.close()


def _reproducible_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed timestamps so the DOCX hash is stable."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_DT.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(DOCX_TITLE, level=1)
    for para in DOCX_PARAGRAPHS:
        d.add_paragraph(para)
    table = d.add_table(rows=0, cols=2)
    for label, value in DOCX_SIGNATURE:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    props = d.core_properties
    props.author = "make_fixtures.py"
    props.title = "Witness statement (fictional)"
    props.created = props.modified = FIXED_DT
    props.last_modified_by = "make_fixtures.py"
    buf = io.BytesIO()
    d.save(buf)
    path.write_bytes(_reproducible_zip(buf.getvalue()))


def make_eml(path: Path) -> None:
    top_from, top_to, top_date, top_subj, top_body = EML_MESSAGES[0]
    parts = [top_body]
    for frm, to, sent, subj, body in EML_MESSAGES[1:]:
        parts.append(
            f"-----Original Message-----\nFrom: {frm}\nSent: {sent}\nTo: {to}\n"
            f"Subject: {subj}\n\n{body}"
        )
    msg = EmailMessage()
    msg["From"] = top_from
    msg["To"] = top_to
    msg["Date"] = top_date
    msg["Subject"] = top_subj
    msg["Message-ID"] = "<demo-0325@halvorsen-freight.example>"
    msg.set_content("\n\n".join(parts) + "\n")
    path.write_bytes(msg.as_bytes(policy=SMTP))  # CRLF line endings on purpose


def make_all(out_dir: Path, gold_dir: Path) -> list[Path]:
    """Write the five fixture documents plus the gold files; returns written paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    gold_dir.mkdir(parents=True, exist_ok=True)
    written = [out_dir / n for n in (PDF_NAME, DOCX_NAME, EML_NAME, TXT_NAME, MD_NAME)]
    make_pdf(written[0])
    make_docx(written[1])
    make_eml(written[2])
    written[3].write_text(TXT_TEXT, encoding="utf-8", newline="\n")
    written[4].write_text(MD_TEXT, encoding="utf-8", newline="\n")
    for name, payload in (
        ("demo_timeline.json", gold_timeline()),
        ("demo_disputes.json", gold_disputes()),
    ):
        p = gold_dir / name
        p.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        written.append(p)
    return written


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=ROOT / "tests" / "fixtures" / "case_demo")
    ap.add_argument("--gold", type=Path, default=ROOT / "eval" / "gold")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    for p in make_all(args.out, args.gold):
        log.info("wrote %s", p)


if __name__ == "__main__":
    main()
